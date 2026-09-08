from __future__ import annotations
import csv, io, json, os, re, uuid
from datetime import date
from pathlib import Path
from typing import Any

import pdfplumber
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
app = FastAPI(title="Quick Filler", version="1.0.0")
app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")

MAX_BYTES = 15 * 1024 * 1024
ALLOWED_TYPES = {"cartao-ponto", "holerite"}
STORE: dict[str, dict[str, Any]] = {}


def validate_pdf(data: bytes, filename: str | None = None) -> None:
    if not data:
        raise HTTPException(400, "Arquivo vazio")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "PDF excede o limite de 15 MB")
    if not data.startswith(b"%PDF-"):
        raise HTTPException(400, "O arquivo enviado não é um PDF válido")
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data), strict=False)
        if len(reader.pages) == 0:
            raise ValueError
    except Exception as exc:
        raise HTTPException(400, "PDF corrompido ou ilegível") from exc


def ocr_page(page) -> str:
    try:
        import pytesseract
        image = page.to_image(resolution=220).original
        return pytesseract.image_to_string(image, lang=os.getenv("TESSERACT_LANG", "por+eng"))
    except Exception:
        return ""


def extract_page_text(page) -> str:
    text = (page.extract_text() or "").strip()
    return text if text else ocr_page(page).strip()


def normalize_time(raw: str) -> str:
    raw = raw.strip().replace(" ", "")
    m = re.fullmatch(r"(\d{1,2})[:hH](\d{2})", raw)
    if not m:
        return "?" if not raw else raw
    h, minute = int(m.group(1)), int(m.group(2))
    if not (0 <= h <= 23 and 0 <= minute <= 59):
        return "?:??"
    return f"{h:02d}:{minute:02d}"


def parse_cartao(text: str, page_no: int) -> dict[str, Any]:
    days = []
    date_re = re.compile(r"(?<!\d)(\d{1,2}/\d{1,2}(?:/\d{2,4})?)(?!\d)")
    time_re = re.compile(r"(?<!\d)(\d{1,2})[:hH](\d{2})(?!\d)")
    for line in text.splitlines():
        dm = date_re.search(line)
        if not dm:
            continue
        raw_date = dm.group(1)
        punches = []
        for idx, match in enumerate(time_re.finditer(line)):
            raw_time = f"{match.group(1)}:{match.group(2)}"
            punches.append({
                "kind": "IN" if idx % 2 == 0 else "OUT",
                "time_raw": raw_time,
                "time_hhmm": normalize_time(raw_time),
            })
        days.append({"date_raw": raw_date, "punches": punches})
    return {"pages": [{"page": page_no, "days": days}]}


BASE_LABEL_RE = re.compile(r"^(base\s|total\s|valor\s+líquido|valor\s+liquido)", re.I)
MONEY_RE = re.compile(r"[\d?]{1,3}(?:\.[\d?]{3})*,[\d?]{2}")


def parse_holerite(text: str, page_no: int) -> dict[str, Any]:
    year_match = re.search(r"\b(20\d{2})\b", text)
    year = year_match.group(1) if year_match else "?"
    month_match = re.search(r"(?:compet[êe]ncia|m[êe]s)\s*[:\-/]?\s*(0?[1-9]|1[0-2])\b", text, re.I)
    month = f"{int(month_match.group(1)):02d}" if month_match else "?"
    fields, bases = [], []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        money_matches = list(MONEY_RE.finditer(line))
        if not money_matches:
            continue
        value = money_matches[-1].group(0)
        prefix = line[:money_matches[-1].start()].strip(" -\t")
        code_match = re.match(r"^(\d{3,5})\s+(.+)$", prefix)
        if code_match:
            code, label_part = code_match.group(1), code_match.group(2).strip()
        else:
            code, label_part = "", prefix
        reference = ""
        ref_match = re.search(r"(?<![\d,])\d+(?:,\d+)?\s*$", label_part)
        if ref_match and not label_part.lower().startswith(("base ", "total ", "valor ")):
            reference = ref_match.group(0)
            label_part = label_part[:ref_match.start()].strip()
        label = label_part.strip(" -")
        if not label:
            continue
        item = {"code": code, "label": label, "reference": reference, "value": value}
        if BASE_LABEL_RE.match(label):
            bases.append({"label": label, "value": value})
        else:
            fields.append(item)
    return {"pages": [{"page": page_no, "year": year, "month": month, "fields": fields, "bases": bases}]}


def process_pdf(data: bytes, kind: str) -> dict[str, Any]:
    pages: list[dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            text = extract_page_text(page)
            parsed = parse_cartao(text, page_no) if kind == "cartao-ponto" else parse_holerite(text, page_no)
            pages.extend(parsed["pages"])
    return {"pages": pages}


def has_uncertainty(obj: Any) -> bool:
    return "?" in json.dumps(obj, ensure_ascii=False)


def parse_date(raw: str) -> date | None:
    if "?" in raw:
        return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", raw.strip())
    if not m:
        return None
    y = int(m.group(3)); y = y if y >= 100 else 2000 + y
    try:
        return date(y, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def warnings_for(value: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    result = []
    pages = value.get("pages", [])
    if kind == "cartao-ponto":
        previous = None
        for p in pages:
            for day in p.get("days", []):
                reasons = []
                punches = day.get("punches", [])
                if len(punches) % 2:
                    reasons.append("Batidas ímpares")
                if has_uncertainty(day):
                    reasons.append("Leitura incerta")
                current = parse_date(day.get("date_raw", ""))
                if current and previous and (current - previous).days != 1:
                    reasons.append("Data não sequencial")
                if current:
                    previous = current
                result.append({"page": p.get("page"), "date_raw": day.get("date_raw", ""), "reasons": reasons})
    else:
        previous = None
        for p in pages:
            reasons = []
            if not p.get("fields") and not p.get("bases"):
                reasons.append("Página vazia")
            if has_uncertainty(p):
                reasons.append("Leitura incerta")
            try:
                month_key = int(p.get("year")) * 12 + int(p.get("month"))
                if previous is not None and month_key != previous + 1:
                    reasons.append("Mês não sequencial")
                if p.get("month") != "?":
                    previous = month_key
            except (TypeError, ValueError):
                pass
            result.append({"page": p.get("page"), "reasons": reasons})
    return result


def row_has_warning(reasons: list[str]) -> str:
    if any("sequencial" in r for r in reasons):
        return "red"
    if reasons:
        return "yellow"
    return ""


def workbook_for(item: dict[str, Any]) -> io.BytesIO:
    wb = Workbook(); ws = wb.active; ws.title = "Transcrição"
    value, kind = item["value"], item["tipo"]
    warning_rows = warnings_for(value, kind)
    warn_by_key = {}
    if kind == "cartao-ponto":
        for w in warning_rows: warn_by_key[(w["page"], w["date_raw"])] = w["reasons"]
        max_punches = max((len(d.get("punches", [])) for p in value.get("pages", []) for d in p.get("days", [])), default=0)
        headers = ["Data"] + [f"{'Entrada' if i % 2 == 0 else 'Saída'} {i // 2 + 1}" for i in range(max_punches)]
        ws.append(headers)
        for p in value.get("pages", []):
            for d in p.get("days", []):
                ws.append([d.get("date_raw", "")] + [x.get("time_hhmm", "") for x in d.get("punches", [])])
                reasons = warn_by_key.get((p.get("page"), d.get("date_raw", "")), [])
                if row_has_warning(reasons):
                    color = "F8D7DA" if row_has_warning(reasons) == "red" else "FFF3CD"
                    for cell in ws[ws.max_row]: cell.fill = PatternFill("solid", fgColor=color)
    else:
        labels = []
        for p in value.get("pages", []):
            for f in p.get("fields", []):
                if f.get("label") not in labels: labels.append(f.get("label"))
        ws.append(["Pág.", "Mês", "Ano"] + labels)
        for p in value.get("pages", []):
            mapping = {f.get("label"): f.get("value", "") for f in p.get("fields", [])}
            ws.append([p.get("page"), p.get("month", ""), p.get("year", "")] + [mapping.get(label, "") for label in labels])
            reasons = next((w["reasons"] for w in warning_rows if w["page"] == p.get("page")), [])
            if row_has_warning(reasons):
                color = "F8D7DA" if row_has_warning(reasons) == "red" else "FFF3CD"
                for cell in ws[ws.max_row]: cell.fill = PatternFill("solid", fgColor=color)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173772")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(12, min(42, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width
    thin_red = Side(style="thin", color="DC3545")
    for row in ws.iter_rows():
        for cell in row:
            cell.number_format = "@"
    for idx, w in enumerate(warning_rows, start=2):
        if row_has_warning(w["reasons"]) == "red" and idx <= ws.max_row:
            ws.cell(idx, 1).border = Border(left=thin_red)
    out = io.BytesIO(); wb.save(out); out.seek(0); return out


def csv_for(item: dict[str, Any]) -> io.BytesIO:
    value, kind = item["value"], item["tipo"]
    out = io.StringIO(); writer = csv.writer(out, lineterminator="\n")
    if kind == "cartao-ponto":
        max_punches = max((len(d.get("punches", [])) for p in value.get("pages", []) for d in p.get("days", [])), default=0)
        writer.writerow(["Data"] + [f"{'Entrada' if i % 2 == 0 else 'Saída'} {i // 2 + 1}" for i in range(max_punches)])
        for p in value.get("pages", []):
            for d in p.get("days", []): writer.writerow([d.get("date_raw", "")] + [x.get("time_hhmm", "") for x in d.get("punches", [])])
    else:
        labels = []
        for p in value.get("pages", []):
            for f in p.get("fields", []):
                if f.get("label") not in labels: labels.append(f.get("label"))
        writer.writerow(["Pág.", "Mês", "Ano"] + labels)
        for p in value.get("pages", []):
            mapping = {f.get("label"): f.get("value", "") for f in p.get("fields", [])}
            writer.writerow([p.get("page"), p.get("month", ""), p.get("year", "")] + [mapping.get(l, "") for l in labels])
    return io.BytesIO(out.getvalue().encode("utf-8-sig"))


@app.get("/")
def index(): return FileResponse(BASE / "templates" / "index.html")

@app.get("/healthz")
def healthz(): return {"status": "ok"}

@app.post("/api/transcricoes")
async def create_transcricao(arquivo: UploadFile = File(...), tipo: str = Form(...)):
    if tipo not in ALLOWED_TYPES: raise HTTPException(400, "tipo deve ser cartao-ponto ou holerite")
    data = await arquivo.read()
    validate_pdf(data, arquivo.filename)
    tid = uuid.uuid4().hex
    STORE[tid] = {"id": tid, "tipo": tipo, "status": "processando", "erro": None, "value": None, "pdf": data}
    try:
        STORE[tid]["value"] = process_pdf(data, tipo)
        STORE[tid]["status"] = "concluido"
    except Exception:
        STORE[tid]["status"] = "erro"
        STORE[tid]["erro"] = "Não foi possível processar o PDF"
    return JSONResponse(status_code=202, content={"id": tid})

@app.get("/api/transcricoes/{tid}")
def get_transcricao(tid: str):
    item = STORE.get(tid)
    if not item: raise HTTPException(404, "transcrição não encontrada")
    return {k: item[k] for k in ("id", "tipo", "status", "erro", "value")}

@app.put("/api/transcricoes/{tid}")
async def update_transcricao(tid: str, payload: dict[str, Any]):
    item = STORE.get(tid)
    if not item: raise HTTPException(404, "transcrição não encontrada")
    if not isinstance(payload.get("value"), dict): raise HTTPException(400, "value obrigatório")
    item["value"] = payload["value"]
    item["status"] = "concluido"
    return {k: item[k] for k in ("id", "tipo", "status", "erro", "value")}

@app.get("/api/transcricoes/{tid}/avisos")
def get_warnings(tid: str):
    item = STORE.get(tid)
    if not item or not item.get("value"): raise HTTPException(404, "transcrição não encontrada")
    return {"warnings": warnings_for(item["value"], item["tipo"])}

@app.get("/api/transcricoes/{tid}/planilha")
def download_sheet(tid: str, formato: str = "xlsx"):
    item = STORE.get(tid)
    if not item or not item.get("value"): raise HTTPException(404, "transcrição não encontrada")
    if formato == "json":
        return JSONResponse(item["value"], headers={"Content-Disposition": "attachment; filename=transcricao.json"})
    if formato == "csv":
        data = csv_for(item)
        return StreamingResponse(data, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=transcricao.csv"})
    if formato != "xlsx": raise HTTPException(400, "formato deve ser xlsx, csv ou json")
    data = workbook_for(item)
    return StreamingResponse(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=transcricao.xlsx"})

@app.get("/api/transcricoes/{tid}/pdf")
def pdf(tid: str):
    item = STORE.get(tid)
    if not item: raise HTTPException(404, "transcrição não encontrada")
    return StreamingResponse(io.BytesIO(item["pdf"]), media_type="application/pdf", headers={"Content-Disposition": "inline; filename=documento.pdf"})
